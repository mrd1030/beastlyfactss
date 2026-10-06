import io, json, os, re, openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

SP = os.path.dirname(os.path.abspath(__file__))
d = json.load(io.open(os.path.join(SP, 'extract.json'), encoding='utf-8'))
OUT = 'PRICE_REDO.xlsx'
STATUS = {}
if os.path.exists(OUT):  # keep status and notes written so far
    old = openpyxl.load_workbook(OUT)['Summary']
    hdr = [c.value for c in old[1]]
    for r in old.iter_rows(min_row=2, values_only=True):
        if r[0]:
            STATUS[r[0]] = (r[hdr.index('Status')], r[hdr.index('Notes')])

HFILL = PatternFill('solid', fgColor='2F5233')
RED = PatternFill('solid', fgColor='FFC7CE')
GREEN = PatternFill('solid', fgColor='C6EFCE')
YEL = PatternFill('solid', fgColor='FFEB9C')
WRAP = Alignment(wrap_text=True, vertical='top')


def sheet(wb, title, headers, widths, rows, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HFILL
        c.alignment = WRAP
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = ws.dimensions
    return ws


# consistency rows still open on prices
cons = []
cw = openpyxl.load_workbook('CONSISTENCY_CHECK.xlsx', read_only=True)
for ws in cw.worksheets[1:]:
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not r[0]:
            continue
        txt = ' '.join(str(x) for x in r if x)
        if r[3] != '✓ Fixed' and ('$' in txt or 'price' in txt.lower()):
            cons.append((ws.title, i + 1) + tuple(r))


def key_for(name):
    k = name.lower().replace("'", '').replace(' ', '-')
    alias = {'african-fat-tailed-gecko': 'african-fat-tail', 'argentine-black-and-white-tegu': 'argentine-tegu', 'betta': 'betta-fish',
             'white-s-tree-frog': 'whites-tree-frog', 'hognose': 'hognose-snake', 'oscar': 'oscar-fish', 'corydoras': 'corydoras-catfish',
             'cory-catfish': 'corydoras-catfish', 'pleco': 'bristlenose-pleco', 'jackson-s-chameleon': 'jacksons-chameleon', 'pac-man-frog': 'pacman-frog'}
    return alias.get(k, k)


animals = sorted(d['animals'], key=lambda a: (a['package'] == 'yes', a['animal']))
keys = {a['animal'] for a in animals}
nm = {}
for m in d['mentions']:
    nm.setdefault(m[0], {}).setdefault(m[1], 0)
    nm[m[0]][m[1]] += 1
ncons = {}
for c in cons:
    k = key_for(c[2])
    ncons[k] = ncons.get(k, 0) + 1

wb = openpyxl.Workbook()
srows = []
for a in animals:
    st, notes = STATUS.get(a['animal'], ('To do', ''))
    srows.append([a['animal'], a['group'], 'package (site only)' if a['package'] else 'full redo', st,
                  a['setup_h2'], a['setup_sum'], a['setup_ok'], a['monthly_h2'], a['monthly_sum'], a['monthly_ok'], a['annual_sum'],
                  a['budget'], a['route'], a.get('hub_miss', ''), a['n_rows'], a['n_unlinked'], a['n_round'], a['n_cover'], a['n_buy_unlinked'],
                  a['n_buy_not_costed'], nm.get(a['animal'], {}).get('other page', 0) + nm.get(a['animal'], {}).get('encyclopedia', 0),
                  ncons.get(a['animal'], 0), notes])
ws = sheet(wb, 'Summary', ['Animal', 'Group', 'Scope', 'Status', 'Setup heading', 'Setup rows add to', 'Heading matches rows',
                           'Monthly heading', 'Monthly rows add to', 'Heading matches rows', 'Annual rows add to', 'Hub Budget row',
                           'Hub cost route line', 'Hub figures missing from cost guide', 'Cost rows', 'Rows with no gear link', 'Rows not in $5 steps', 'Rows below catalog price',
                           'Hub buy list items with no link', 'Hub items with no cost row', 'Other pages pricing it', 'Open consistency rows', 'Notes'],
           [20, 11, 14, 10, 34, 18, 12, 30, 18, 12, 18, 50, 50, 18, 8, 9, 9, 9, 9, 9, 9, 9, 50], srows, first=True)
for row in ws.iter_rows(min_row=2):
    for c in (row[6], row[9]):
        c.fill = GREEN if c.value == 'yes' else (RED if c.value == 'NO' else PatternFill())
    row[13].fill = RED if row[13].value else GREEN
    for c in row[15:22]:
        if isinstance(c.value, int) and c.value:
            c.fill = YEL
    row[3].fill = GREEN if row[3].value == 'Done' else (YEL if row[3].value == 'In progress' else PatternFill())

order = {a['animal']: i for i, a in enumerate(animals)}
rr = sorted(d['rows'], key=lambda r: order.get(r['animal'], 999))
sheet(wb, 'Cost rows', ['Animal', 'Table heading', 'Kind', 'Item', 'Price as written', 'Low', 'High', 'Linked product (catalog price)',
                        'No gear link', 'Below catalog price', 'Rounding', 'Total row'],
      [18, 34, 9, 50, 22, 8, 8, 50, 10, 22, 14, 10],
      [[r['animal'], r['table'], r['kind'], r['item'], r['price'], r['lo'], r['hi'], r['linked'], r['unlinked'], r['covers'], r['rounding'], r.get('is_total', '')] for r in rr])
bb = sorted(d['buy'], key=lambda b: order.get(b['animal'], 999))
sheet(wb, 'Hub buy list', ['Animal', 'Buy list item', 'Links to', 'Catalog price', 'Has a cost row'], [18, 60, 40, 16, 12],
      [[b['animal'], b['item'], b['product'] or '(no link)', b['price'], b['in_cost']] for b in bb])
mm = sorted(d['mentions'], key=lambda m: (order.get(m[0], 999), ['cost guide', 'hub', 'other page', 'encyclopedia'].index(m[1])))
sheet(wb, 'Price mentions', ['Animal', 'Where', 'File', 'Field', 'Text'], [18, 12, 34, 16, 110], mm)
sheet(wb, 'Consistency rows', ['Workbook tab', 'Row', 'Animal', 'Guide', 'File', 'Consistent', 'What is inconsistent', 'Conflicts with',
                               'Severity', 'Proposed fix', 'Notes'],
      [12, 6, 20, 16, 30, 8, 50, 40, 8, 40, 60], cons)
wb.save(OUT)
print('saved', OUT, len(srows), 'animals', len(rr), 'rows', len(bb), 'buy', len(mm), 'mentions', len(cons), 'consistency rows')
